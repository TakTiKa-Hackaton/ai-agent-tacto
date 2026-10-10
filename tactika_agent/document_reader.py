"""Attachment text extraction and quote-term recognition for supplier documents."""

from __future__ import annotations

import io
import re

FIELD_LABELS = {
    "Unit price (EUR)": "unit_price",
    "Payment terms (days net)": "payment_days",
    "Lead time (days)": "lead_time_days",
    "Volume per year (pcs)": "volume_per_year",
    "Contract term (months)": "contract_months",
    "Price valid from": "effective_date",
    "Index clause": "index_clause",
    "Incoterm": "incoterm",
    "Tooling (EUR)": "tooling_eur",
    "Quality commitment (ppm)": "ppm_commitment",
    "Chargeback per complaint (EUR)": "chargeback_eur_per_complaint",
}
HEADINGS = {
    "Offered terms": "stated",
    "Our standard terms": "standard",
    "Terms fixed by your enquiry": "fixed",
    "Agreed terms": "agreed",
}
INDEX_PATTERN = re.compile(r"(\S+), (\d+)% material, deadband ([\d.]+) %, base (\S+)")


def extract_text(filename, content):
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    if lower_name.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), data_only=False)
        return "\n".join(
            f"[{sheet.title}]\n" + "\n".join(
                " | ".join("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            )
            for sheet in workbook.worksheets
        )
    return content.decode("utf-8", "replace")


def _parse_value(field, raw):
    if field == "index_clause":
        match = INDEX_PATTERN.fullmatch(raw)
        if not match:
            return None
        series, share, band, base = match.groups()
        return {"series": series, "material_share": int(share) / 100,
                "deadband_pct": float(band), "base_month": base}
    if field in {"effective_date", "incoterm"}:
        return raw
    number = float(raw.replace(",", "."))
    return number if field in {"unit_price", "tooling_eur", "chargeback_eur_per_complaint"} else int(number)


def parse_quote(text):
    result, section = {}, None
    for line in text.splitlines():
        line = line.strip()
        for heading, name in HEADINGS.items():
            if line.startswith(heading):
                section = name
                break
        for label, field in FIELD_LABELS.items():
            if section and line.startswith(label + " "):
                try:
                    result.setdefault(section, {})[field] = _parse_value(field, line[len(label):].strip())
                except (ValueError, TypeError):
                    pass
                break
    return result


def supplier_standard_terms(events):
    for message in reversed(events):
        for attachment in message.get("attachments") or []:
            name = attachment.get("name", "")
            if name.startswith("quote") and attachment.get("text"):
                return parse_quote(attachment["text"]).get("standard", {})
    return {}
